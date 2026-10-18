import csv
import hmac
import io
import json
from io import BytesIO
from uuid import uuid4

import streamlit as st

from src.config import get_settings
from src.db import (
    create_review_run,
    get_chunks_for_citations,
    get_paper_chunks,
    get_paper_id,
    get_run_notes,
    init_db,
    replace_paper_chunks,
    save_citation_check,
    save_paper_note,
    save_papers,
    update_paper_download,
    update_paper_status,
)
from src.llm.client import LLMError, generate_keyword_suggestion
from src.notes.generator import generate_paper_note
from src.notes.parser import parse_pdf
from src.papers.download import download_open_access_pdf
from src.papers.keys import build_citation_keys
from src.review.checker import check_claim, extract_claims
from src.review.generator import generate_review
from src.search.openalex import OpenAlexError, search_works


st.set_page_config(page_title="文献综述 Agent", page_icon="📚", layout="wide")
settings = get_settings()

st.markdown("""
<style>
.hero {padding: 2.2rem 2.5rem; border-radius: 24px; background: linear-gradient(135deg,#312e81,#4f46e5 55%,#818cf8); color:white; margin-bottom:1.5rem;}
.hero h1 {font-size:2.7rem; margin:0 0 .5rem 0; color:white;}
.hero p {font-size:1.05rem; opacity:.9; margin:0;}
.step-card {padding:1rem 1.2rem; border:1px solid #e5e7eb; border-radius:16px; background:#fff; margin-bottom:1rem;}
.muted {color:#6b7280;}
</style>
""", unsafe_allow_html=True)

if settings.app_password:
    st.markdown('<div class="hero"><h1>文献综述 Agent</h1><p>可追溯的学术检索、精读与引文核对工作流</p></div>', unsafe_allow_html=True)
    password = st.text_input("请输入访问密码", type="password")
    if not password:
        st.info("该 Demo 已设置访问密码。")
        st.stop()
    if not hmac.compare_digest(password, settings.app_password):
        st.error("访问密码不正确。")
        st.stop()

init_db()

steps = ["主题与关键词", "检索与选文", "PDF 获取", "精读笔记", "综述与核对"]
if "current_step" not in st.session_state:
    st.session_state.current_step = 1
current_step = st.session_state.current_step

st.markdown('<div class="hero"><h1>文献综述 Agent</h1><p>从研究主题到可追溯综述，一步步完成检索、精读、生成与核对。</p></div>', unsafe_allow_html=True)

with st.sidebar:
    st.markdown("### 工作流进度")
    for i, label in enumerate(steps, 1):
        icon = "✅" if i < current_step else ("🔵" if i == current_step else "⚪")
        st.write(f"{icon} **{i}. {label}**")
    st.progress((current_step - 1) / len(steps))
    st.divider()
    st.caption("SciencePRO Demo · DeepSeek + OpenAlex")

def go(step: int):
    st.session_state.current_step = step
    st.rerun()

if current_step == 1:
    st.header("1. 确认研究主题")
    st.caption("研究主题必填，年份范围可选。")
    topic = st.text_area("研究主题 *", value=st.session_state.get("topic", ""), placeholder="例如：钙钛矿太阳能电池稳定性提升策略", height=110)
    c1, c2 = st.columns(2)
    with c1:
        use_year = st.checkbox("限定年份范围", value=bool(st.session_state.get("year_range")))
    with c2:
        if use_year:
            year_from = st.number_input("起始年份", 1900, 2100, 2018)
            year_to = st.number_input("结束年份", 1900, 2100, 2025)
        else:
            year_from, year_to = 1900, 2100
    if st.button("生成关键词与综述大纲 →", type="primary", disabled=not topic.strip()):
        run_id = st.session_state.get("run_id", uuid4().hex[:12])
        st.session_state.update(topic=topic.strip(), year_range=(year_from, year_to), run_id=run_id)
        create_review_run(run_id, topic.strip())
        with st.spinner("正在生成关键词和综述大纲..."):
            try:
                st.session_state.keyword_suggestion = generate_keyword_suggestion(topic.strip(), year_from, year_to)
                go(2)
            except LLMError as exc:
                st.error(str(exc))

elif current_step == 2:
    suggestion = st.session_state.get("keyword_suggestion", {})
    st.header("2. 关键词与检索范围")
    title = st.text_input("综述标题", suggestion.get("title", ""))
    zh = st.text_input("中文关键词（逗号分隔）", "，".join(suggestion.get("keywords_zh", [])))
    en = st.text_input("英文关键词（逗号分隔）", ", ".join(suggestion.get("keywords_en", [])))
    st.markdown("#### 建议综述大纲")
    outline = suggestion.get("outline", [])
    for i, item in enumerate(outline, 1):
        st.write(f"{i}. {item}")
    st.code(suggestion.get("search_query", ""), language="text")
    c1, c2 = st.columns(2)
    with c1:
        if st.button("← 返回修改主题"):
            go(1)
    with c2:
        if st.button("确认关键词，开始检索 →", type="primary"):
            st.session_state.confirmed_keywords = {"title": title, "keywords_zh": [x.strip() for x in zh.replace("，", ",").split(",") if x.strip()], "keywords_en": [x.strip() for x in en.split(",") if x.strip()], "search_query": suggestion.get("search_query", ""), "outline": outline}
            with st.spinner("正在检索 OpenAlex..."):
                try:
                    query = " AND ".join(st.session_state.confirmed_keywords["keywords_en"][:4])
                    count, papers = search_works(query, *st.session_state.year_range, per_page=10)
                    st.session_state.openalex_result = {"hit_count": count, "papers": [p.to_dict() for p in papers]}
                    go(3)
                except OpenAlexError as exc:
                    st.error(str(exc))
    if st.button("← 返回主题设置", key="back_step_2"):
        go(1)

elif current_step == 3:
    result = st.session_state.get("openalex_result", {})
    papers = result.get("papers", [])
    st.header("3. 选择文献与获取 PDF")
    st.success(f"OpenAlex 命中约 {result.get('hit_count', 0):,} 篇，当前展示 {len(papers)} 篇。")
    st.caption("默认选择前 5 篇，最多进入精读流程 10 篇。")
    selected = []
    for i, paper in enumerate(papers):
        checked = st.checkbox(f"{i+1}. {paper['title']}", value=i < 5, key=f"pick_{st.session_state.run_id}_{i}")
        if checked:
            selected.append(i)
        with st.expander("查看元数据"):
            st.write(f"作者：{', '.join(paper['authors'][:5]) or '未知'}")
            st.write(f"年份：{paper['year'] or '未知'} · DOI：{paper['doi'] or '无'} · 被引：{paper['cited_by_count']}")
    if len(selected) > 10:
        st.warning("最多选择 10 篇。")
    if st.button("确认文献并获取 PDF →", type="primary", disabled=not selected or len(selected) > 10):
        chosen = [papers[i] for i in selected]
        for paper, key in zip(chosen, build_citation_keys(chosen)):
            paper["citation_key"] = key
        save_papers(st.session_state.run_id, chosen)
        st.session_state.selected_papers = chosen
        with st.spinner("正在查找开放获取 PDF..."):
            results = []
            for paper in chosen:
                r = download_open_access_pdf(paper)
                paper["pdf_path"], paper["status"] = r.pdf_path, r.status
                update_paper_download(st.session_state.run_id, r.citation_key, r.status, r.pdf_path, r.source_url)
                results.append(r)
            st.session_state.download_results = results
        go(4)
    if st.button("← 返回关键词设置", key="back_step_3"):
        go(2)

elif current_step == 4:
    st.header("4. PDF 获取与精读准备")
    st.caption("自动获取失败的论文可以打开来源链接，或从浏览器/Zotero 上传 PDF。")
    results = st.session_state.get("download_results", [])
    by_key = {r.citation_key: r for r in results}
    failed = []
    for paper in st.session_state.selected_papers:
        r = by_key.get(paper["citation_key"])
        if r and r.status == "pdf_found":
            st.success(f"{paper['citation_key']}：PDF 已获取")
        else:
            failed.append(paper)
            st.warning(f"{paper['citation_key']}：{r.message if r else '待处理'}")
            if r and r.source_url:
                st.link_button("打开来源链接", r.source_url)
            uploaded = st.file_uploader(f"上传 {paper['citation_key']} PDF", type=["pdf"], key=f"upload_{paper['citation_key']}")
            if uploaded:
                target = settings.data_path / "pdfs" / f"{paper['citation_key']}_manual.pdf"
                target.write_bytes(uploaded.getvalue())
                paper["pdf_path"], paper["status"] = str(target), "pdf_found"
                update_paper_download(st.session_state.run_id, paper['citation_key'], 'pdf_found', str(target), r.source_url if r else None)
                st.success("已上传，可继续解析。")
    ready = [p for p in st.session_state.selected_papers if p.get("pdf_path")]
    if ready and st.button("解析 PDF，生成精读笔记 →", type="primary"):
        parse_results = []
        for paper in ready:
            try:
                chunks = parse_pdf(paper["pdf_path"])
                paper_id = get_paper_id(st.session_state.run_id, paper["citation_key"])
                replace_paper_chunks(paper_id, [c.__dict__ for c in chunks])
                update_paper_status(st.session_state.run_id, paper["citation_key"], "parsed")
                parse_results.append({"citation_key": paper["citation_key"], "status": "parsed", "chunk_count": len(chunks)})
            except Exception as exc:
                parse_results.append({"citation_key": paper["citation_key"], "status": "failed", "message": str(exc)})
        st.session_state.parse_results = parse_results
        go(5)
    if st.button("← 返回文献选择", key="back_step_4"):
        go(3)

elif current_step == 5:
    st.header("5. 生成综述与引文核对")
    parsed = {r["citation_key"] for r in st.session_state.get("parse_results", []) if r["status"] == "parsed"}
    papers = [p for p in st.session_state.selected_papers if p["citation_key"] in parsed]
    st.info(f"已解析 {len(papers)} 篇全文文献。")
    if papers and not st.session_state.get("note_results") and st.button("生成逐篇精读笔记", type="primary"):
        notes = []
        for paper in papers:
            try:
                paper_id = get_paper_id(st.session_state.run_id, paper["citation_key"])
                note = generate_paper_note(paper, get_paper_chunks(paper_id))
                save_paper_note(paper_id, json.dumps(note, ensure_ascii=False), note.get("source_quality", "full_text"), note.get("model_name", ""))
                notes.append({"citation_key": paper["citation_key"], "status": "generated", "note": note})
            except Exception as exc:
                notes.append({"citation_key": paper["citation_key"], "status": "failed", "message": str(exc)})
        st.session_state.note_results = notes
        st.rerun()
    if st.session_state.get("note_results"):
        st.success("精读笔记已生成。")
        generated_note_items = [
            item for item in st.session_state.note_results if item.get("status") == "generated"
        ]
        if generated_note_items:
            notes_json = json.dumps(generated_note_items, ensure_ascii=False, indent=2)
            markdown_parts = ["# 文献精读笔记\n"]
            for item in generated_note_items:
                note = item["note"]
                markdown_parts.append(f"## {item['citation_key']}\n")
                markdown_parts.append(f"### 研究问题\n{note['research_question']['text']}\n")
                markdown_parts.append(f"### 数据与样本\n{note['data_and_sample']['text']}\n")
                markdown_parts.append(f"### 研究方法\n{note['method']['text']}\n")
                markdown_parts.append("### 主要发现\n")
                for finding in note.get("findings", []):
                    markdown_parts.append(f"- {finding['text']}\n")
                    for evidence in finding.get("evidence", []):
                        markdown_parts.append(
                            f"  - 原文证据（第 {evidence.get('page')} 页）：{evidence.get('quote')}\n"
                        )
                markdown_parts.append(f"### 局限性\n{note['limitations']['text']}\n")
            notes_markdown = "\n".join(markdown_parts)

            excel_buffer = BytesIO()
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill

            workbook = Workbook()
            overview = workbook.active
            overview.title = "文献概览"
            overview_headers = ["引用Key", "研究问题", "数据与样本", "方法", "局限性", "来源质量"]
            overview.append(overview_headers)
            for cell in overview[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="4F46E5")
            evidence_sheet = workbook.create_sheet("证据句")
            evidence_sheet.append(["引用Key", "发现", "原文证据", "页码"])
            for cell in evidence_sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="4F46E5")
            for item in generated_note_items:
                note = item["note"]
                overview.append([
                    item["citation_key"],
                    note["research_question"]["text"],
                    note["data_and_sample"]["text"],
                    note["method"]["text"],
                    note["limitations"]["text"],
                    note.get("source_quality", "full_text"),
                ])
                for finding in note.get("findings", []):
                    for evidence in finding.get("evidence", []):
                        evidence_sheet.append([
                            item["citation_key"], finding["text"], evidence.get("quote", ""), evidence.get("page", "")
                        ])
            for sheet in (overview, evidence_sheet):
                sheet.freeze_panes = "A2"
                for column in sheet.columns:
                    sheet.column_dimensions[column[0].column_letter].width = 28
            workbook.save(excel_buffer)

            st.markdown("#### 下载精读笔记")
            d1, d2, d3 = st.columns(3)
            with d1:
                st.download_button("下载 JSON", notes_json, "paper_notes.json", "application/json", key="download_notes_json")
            with d2:
                st.download_button("下载 Markdown", notes_markdown, "paper_notes.md", "text/markdown", key="download_notes_md")
            with d3:
                st.download_button("下载 Excel", excel_buffer.getvalue(), "paper_notes.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="download_notes_excel")
        if not st.session_state.get("review_result") and st.button("生成综述 →", type="primary"):
            try:
                from src.db import save_review
                markdown = generate_review(get_run_notes(st.session_state.run_id), st.session_state.topic, st.session_state.confirmed_keywords["title"])
                rid = save_review(st.session_state.run_id, markdown, markdown.count("[@"))
                st.session_state.review_result = {"id": rid, "markdown": markdown, "citation_count": markdown.count("[@")}
                st.rerun()
            except Exception as exc:
                st.error(str(exc))
    if st.session_state.get("review_result"):
        review = st.session_state.review_result
        st.divider(); st.subheader("综述结果"); st.markdown(review["markdown"])
        docx_buffer = BytesIO()
        from docx import Document
        from docx.shared import Pt

        document = Document()
        for line in review["markdown"].splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith("### "):
                document.add_heading(stripped[4:], level=3)
            elif stripped.startswith("## "):
                document.add_heading(stripped[3:], level=2)
            elif stripped.startswith("# "):
                document.add_heading(stripped[2:], level=1)
            elif stripped.startswith("- "):
                document.add_paragraph(stripped[2:], style="List Bullet")
            else:
                document.add_paragraph(stripped)
        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Microsoft YaHei"
                run.font.size = Pt(10.5)
        document.save(docx_buffer)
        d1, d2 = st.columns(2)
        with d1:
            st.download_button("下载 Markdown", review["markdown"], "literature_review.md", "text/markdown", key="download_review_md")
        with d2:
            st.download_button("下载 Word", docx_buffer.getvalue(), "literature_review.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", key="download_review_docx")
        if not st.session_state.get("citation_checks") and st.button("开始引文核对", type="primary"):
            checks = []
            for claim in extract_claims(review["markdown"]):
                try:
                    result = check_claim(claim["claim_text"], claim["citation_keys"], get_chunks_for_citations(st.session_state.run_id, claim["citation_keys"]))
                    save_citation_check(review["id"], claim["claim_text"], claim["citation_keys"], result)
                    checks.append({**claim, **result})
                except Exception as exc:
                    checks.append({**claim, "verdict": "needs_human_review", "evidence": [], "confidence": 0, "reason": str(exc)})
            st.session_state.citation_checks = checks
            st.rerun()
    if st.session_state.get("citation_checks"):
        checks = st.session_state.citation_checks
        labels = {"supported": "支持", "partially_supported": "部分支持", "unsupported": "不支持", "needs_human_review": "需人工复核"}
        counts = {key: sum(x.get("verdict") == key for x in checks) for key in labels}
        # 支持率分母排除“需人工复核”，避免全文缺失导致质量指标失真。
        measurable_total = counts["supported"] + counts["partially_supported"] + counts["unsupported"]
        total = len(checks) or 1
        c1, c2, c3, c4 = st.columns(4)
        support_rate = counts["supported"] / measurable_total if measurable_total else 0
        c1.metric("引文支持率", f"{support_rate:.1%}")
        c2.metric("完全支持", counts["supported"])
        c3.metric("部分支持", counts["partially_supported"])
        c4.metric("不支持", counts["unsupported"])
        st.metric("需人工复核", counts["needs_human_review"])
        st.caption("引文支持率 = 完全支持数 ÷（完全支持数 + 部分支持数 + 不支持数）；需人工复核不计入分母。")
        for i, item in enumerate(checks, 1):
            with st.expander(f"{i}. {labels.get(item.get('verdict'), item.get('verdict'))} · {item['claim_text']}"):
                st.write("引用：", ", ".join(item["citation_keys"]))
                st.write(item.get("reason", ""))
                for evidence in item.get("evidence", []):
                    st.caption(f"第 {evidence.get('page')} 页：{evidence.get('quote')}")
        st.success("流程完成：综述、引文核对与质量指标均已生成。")
    if st.button("← 返回精读笔记", key="back_step_5"):
        go(4)
